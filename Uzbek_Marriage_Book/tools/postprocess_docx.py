#!/usr/bin/env python3
"""Final touches on the pandoc DOCX: core properties and a page break after the title block."""
import sys
from docx import Document
from docx.enum.text import WD_BREAK

path = sys.argv[1]
doc = Document(path)
cp = doc.core_properties
cp.title = "Bir umrlik nikoh"
cp.subject = "Mehr-muhabbat, rahm-shafqat va baxtli oila qurish sanʼati"
cp.language = "uz-Latn-UZ"
cp.keywords = "nikoh; oila; Islom; psixologiya; oʻzbek tili"
cp.author = "Bir umrlik nikoh loyihasi"

for p in doc.paragraphs:
    if p.style.name in ("Date", "Author"):
        p.add_run().add_break(WD_BREAK.PAGE)
        break
doc.save(path)
print("postprocessed", path)
